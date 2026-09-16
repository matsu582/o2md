"""x2md の表検出ロジックの単体テスト

openpyxl でメモリ上に Workbook を組み、一時 xlsx から
ExcelToMarkdownConverter を生成して各判定メソッドを直接呼び出す。
アサーションは現行実装の挙動を固定したもの。
"""

import tempfile
from pathlib import Path

import openpyxl
import pytest
from openpyxl.styles import Border, Side

from o2md.x2md import ExcelToMarkdownConverter

THIN = Side(style="thin")
BORDER_FULL = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BORDER_RIGHT_BOTTOM = Border(right=THIN, bottom=THIN)


def _make_converter(grid, border=None):
    """2次元リスト grid から一時 xlsx を作り converter とシートを返す"""
    wb = openpyxl.Workbook()
    ws = wb.active
    for r, row_vals in enumerate(grid, 1):
        for c, value in enumerate(row_vals, 1):
            cell = ws.cell(row=r, column=c, value=value)
            if border is not None:
                cell.border = border
    tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    wb.save(tmp.name)
    output_dir = tempfile.mkdtemp()
    converter = ExcelToMarkdownConverter(tmp.name, output_dir=output_dir)
    return converter, converter.workbook.active


def _grid(rows, cols, fmt="v{r}{c}"):
    return [[fmt.format(r=r, c=c) for c in range(1, cols + 1)] for r in range(1, rows + 1)]


class TestDetectBorderedTables:
    """_detect_bordered_tables の罫線矩形検出"""

    def test_all_thin_borders_detected_as_one_table(self):
        """4x3 セルを全辺 thin 罫線で囲んだ表 → 1件・範囲一致"""
        converter, ws = _make_converter(_grid(4, 3), border=BORDER_FULL)
        result = converter._detect_bordered_tables(ws, 1, 4, 1, 3)
        assert result == [(1, 4, 1, 3)]

    def test_no_borders_detects_nothing(self):
        """罫線なし → 0件"""
        converter, ws = _make_converter(_grid(4, 3), border=None)
        result = converter._detect_bordered_tables(ws, 1, 4, 1, 3)
        assert result == []

    def test_right_bottom_only_borders_detects_nothing(self):
        """右・下罫線のみ（左・上なし）→ 現行では 0件"""
        converter, ws = _make_converter(_grid(4, 3), border=BORDER_RIGHT_BOTTOM)
        result = converter._detect_bordered_tables(ws, 1, 4, 1, 3)
        assert result == []


class TestIsPlainTextRegion:
    """_is_plain_text_region の表/プレーンテキスト判定"""

    def test_url_like_column_is_not_plain_text(self):
        """13列x6行・全セル埋まり・1列に '/' 含む値 → False（表扱い）"""
        grid = []
        for r in range(6):
            row = []
            for c in range(13):
                row.append("Terraform / Proxmox" if c == 5 else f"item{r}_{c}")
            grid.append(row)
        converter, ws = _make_converter(grid)
        assert converter._is_plain_text_region(ws, (1, 6, 1, 13)) is False

    def test_numbered_list_is_plain_text(self):
        """2列x5行の「1. 2. ...」番号 + 長い右列 → True"""
        grid = [
            [f"{i}.", "これは十五文字以上の説明文テキストです"]
            for i in range(1, 6)
        ]
        converter, ws = _make_converter(grid)
        assert converter._is_plain_text_region(ws, (1, 5, 1, 2)) is True

    def test_single_column_long_text_is_plain_text(self):
        """1列x3行の長文（70字超）→ True"""
        grid = [["あ" * 80] for _ in range(3)]
        converter, ws = _make_converter(grid)
        assert converter._is_plain_text_region(ws, (1, 3, 1, 1)) is True

    def test_normal_table_is_not_plain_text(self):
        """3列x4行の通常表 → False"""
        grid = _grid(4, 3, fmt="d{r}_{c}")
        converter, ws = _make_converter(grid)
        assert converter._is_plain_text_region(ws, (1, 4, 1, 3)) is False


class TestImplicitTableOutput:
    """罫線のみの領域が Markdown テーブルとして出力されること"""

    def test_thirteen_column_thin_border_table_becomes_markdown_table(self, tmp_path):
        """13列 thin 罫線の表が | 区切りと13列分の --- を持つ Markdown 表になる"""
        wb = openpyxl.Workbook()
        ws = wb.active
        for r in range(1, 6):
            for c in range(1, 14):
                cell = ws.cell(row=r, column=c, value=f"c{r}_{c}")
                cell.border = BORDER_FULL
        xlsx = tmp_path / "implicit_13col.xlsx"
        wb.save(xlsx)

        converter = ExcelToMarkdownConverter(str(xlsx), output_dir=str(tmp_path / "out"))
        md_text = Path(converter.convert()).read_text(encoding="utf-8")

        assert "| c1_1 " in md_text
        separator_lines = [l for l in md_text.splitlines() if l.strip().startswith("|") and "---" in l]
        assert separator_lines, "Markdown テーブルの区切り行がありません"
        assert separator_lines[0].count("---") == 13
