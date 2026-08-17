"""MSPDIとWord拡張変換機能の回帰テスト。"""

import tempfile
import zipfile
import io
import sys
import xml.etree.ElementTree as ET
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from o2md.d2md_fields import convert_paragraph, normalize_markdown_url
from o2md.d2md_notes import NoteManager
from o2md.d2md_numbering import NumberingResolver
from o2md.d2md_tables import render_table
from o2md.d2md import WordToMarkdownConverter
from o2md.d2md_composite import (
    absolute_composite_paragraph_xml,
    absolute_drawing_xml,
    composite_paragraphs_xml,
    paragraph_properties_xml,
    section_properties_xml,
)
from o2md.filter import detect_type_from_bytes
import o2md.filter as filter_cli
import o2md.i18n as i18n
from o2md.mspdi import is_mspdi_xml
from o2md.mpp2md import resources_to_markdown_table
from o2md.o2md import detect_file_type
from o2md.xml_safe import fromstring as safe_fromstring
from o2md.x2md import ExcelToMarkdownConverter


W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
PIC = "http://schemas.openxmlformats.org/drawingml/2006/picture"
WPS = "http://schemas.microsoft.com/office/word/2010/wordprocessingShape"


def _drawing_with_kind(kind, width=100, height=100):
    drawing = OxmlElement("w:drawing")
    container = OxmlElement("wp:anchor" if kind == "wsp" else "wp:inline")
    extent = OxmlElement("wp:extent")
    extent.set("cx", str(width))
    extent.set("cy", str(height))
    container.append(extent)
    if kind == "wsp":
        container.append(etree.Element(f"{{{WPS}}}wsp"))
    else:
        pic = etree.Element(f"{{{PIC}}}pic")
        pic.append(etree.Element(f"{{{PIC}}}blipFill"))
        container.append(pic)
    drawing.append(container)
    return drawing


def _fake_composite_converter(document, calls):
    converter = object.__new__(WordToMarkdownConverter)
    converter.doc = document
    converter._composite_skip_paragraphs = set()
    converter._process_mixed_drawings_as_vector = lambda drawings, texts, paragraphs=None: (
        calls.append((drawings, texts, paragraphs)) or True
    )
    return converter


def test_word_composite_figure_keeps_picture_and_shape_in_one_render():
    document = Document()
    paragraph = document.add_paragraph()
    paragraph._p.append(_drawing_with_kind("wsp"))
    paragraph._p.append(_drawing_with_kind("pic"))
    calls = []
    converter = _fake_composite_converter(document, calls)

    assert converter._process_composite_figure(paragraph) is True
    assert len(calls) == 1
    assert len(calls[0][0]) == 2
    assert calls[0][0][0].xpath('.//*[local-name()="wsp"]')
    assert calls[0][0][1].xpath('.//*[local-name()="pic"]')
    assert calls[0][2] == [paragraph]


def test_word_composite_figure_merges_safe_adjacent_shape_and_picture():
    document = Document()
    shape_paragraph = document.add_paragraph()
    shape_paragraph._p.append(_drawing_with_kind("wsp", 100, 100))
    picture_paragraph = document.add_paragraph()
    picture_paragraph._p.append(_drawing_with_kind("pic", 200, 200))
    calls = []
    converter = _fake_composite_converter(document, calls)

    assert converter._process_composite_figure(shape_paragraph) is True
    assert len(calls) == 1
    assert len(calls[0][0]) == 2
    assert [paragraph._p for paragraph in calls[0][2]] == [
        shape_paragraph._p,
        picture_paragraph._p,
    ]
    assert picture_paragraph._p in converter._composite_skip_paragraphs


def test_word_composite_preserves_paragraph_and_section_coordinates():
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.alignment = 1
    drawings_xml = composite_paragraphs_xml(
        [paragraph], [["<w:drawing/>"]]
    )
    assert 'w:jc w:val="center"' in drawings_xml
    assert paragraph_properties_xml(paragraph) in drawings_xml
    assert "w:pgSz" in section_properties_xml(document)
    assert "w:pgMar" in section_properties_xml(document)


def test_absolute_composite_converts_inline_image_and_preserves_extent():
    drawing = f"""
      <w:drawing xmlns:w="{W}"
          xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing">
        <wp:inline>
          <wp:extent cx="1234" cy="5678"/>
          <wp:docPr id="7" name="画像"/>
        </wp:inline>
      </w:drawing>
    """

    converted = absolute_drawing_xml([drawing])
    assert converted is not None
    root = etree.fromstring(converted[0].encode())
    anchor = root.find(f".//{{{WP}}}anchor")
    assert anchor is not None
    assert anchor.get("relativeFrom") is None
    assert anchor.get("behindDoc") == "1"
    assert anchor.find(f"{{{WP}}}extent").get("cx") == "1234"
    assert anchor.find(f"{{{WP}}}extent").get("cy") == "5678"
    assert anchor.find(f".//{{{WP}}}docPr").get("id") == "1"
    assert anchor.find(f"{{{WP}}}positionH").get("relativeFrom") == "margin"
    assert anchor.find(f"{{{WP}}}positionV").find(f"{{{WP}}}posOffset").text == "0"


def test_absolute_composite_rebases_paragraph_anchor_to_margin():
    drawing = f"""
      <w:drawing xmlns:w="{W}"
          xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing">
        <wp:anchor>
          <wp:positionH relativeFrom="column"><wp:posOffset>111</wp:posOffset></wp:positionH>
          <wp:positionV relativeFrom="paragraph"><wp:posOffset>222</wp:posOffset></wp:positionV>
          <wp:extent cx="1234" cy="5678"/>
          <wp:wrapNone/>
        </wp:anchor>
      </w:drawing>
    """

    converted = absolute_drawing_xml([drawing])
    assert converted is not None
    root = etree.fromstring(converted[0].encode())
    anchor = root.find(f".//{{{WP}}}anchor")
    assert anchor.find(f"{{{WP}}}positionH").get("relativeFrom") == "margin"
    assert anchor.find(f"{{{WP}}}positionV").get("relativeFrom") == "margin"
    assert anchor.find(f"{{{WP}}}positionH/{{{WP}}}posOffset").text == "111"
    assert anchor.find(f"{{{WP}}}positionV/{{{WP}}}posOffset").text == "222"


def test_absolute_composite_falls_back_when_position_is_incomplete():
    drawing = f"""
      <w:drawing xmlns:w="{W}"
          xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing">
        <wp:anchor>
          <wp:positionH relativeFrom="column"/>
          <wp:positionV relativeFrom="paragraph"><wp:posOffset>222</wp:posOffset></wp:positionV>
          <wp:extent cx="1234" cy="5678"/>
        </wp:anchor>
      </w:drawing>
    """

    assert absolute_drawing_xml([drawing]) is None


def test_absolute_composite_paragraph_has_minimal_fixed_line_height():
    paragraph_xml = absolute_composite_paragraph_xml(["<w:drawing/>"])
    assert 'w:line="1"' in paragraph_xml
    assert 'w:lineRule="exact"' in paragraph_xml
    assert 'w:sz w:val="1"' in paragraph_xml


def _layout_section_xml(doc_grid=""):
    return f"""
      <w:sectPr xmlns:w="{W}">
        <w:pgSz w:w="1000" w:h="2000"/>
        <w:pgMar w:left="100" w:right="100"/>
        {doc_grid}
      </w:sectPr>
    """


def _inline_drawing_xml(width=254000, height=1000):
    return f"""
      <w:drawing xmlns:w="{W}"
          xmlns:wp="{WP}">
        <wp:inline>
          <wp:extent cx="{width}" cy="{height}"/>
          <wp:docPr id="7" name="画像"/>
        </wp:inline>
      </w:drawing>
    """


def _shape_drawing_xml(vertical_offset=300000):
    return f"""
      <w:drawing xmlns:w="{W}" xmlns:wp="{WP}">
        <wp:anchor>
          <wp:positionH relativeFrom="column"><wp:posOffset>20</wp:posOffset></wp:positionH>
          <wp:positionV relativeFrom="paragraph"><wp:posOffset>{vertical_offset}</wp:posOffset></wp:positionV>
          <wp:extent cx="100" cy="100"/>
        </wp:anchor>
      </w:drawing>
    """


@pytest.mark.parametrize(
    ("jc", "expected_twips"),
    [
        ("center", 200),
        ("distribute", 200),
        ("right", 400),
        ("end", 400),
        (None, 0),
        ("left", 0),
    ],
)
def test_absolute_composite_horizontal_offset_follows_alignment(jc, expected_twips):
    jc_xml = f'<w:jc w:val="{jc}"/>' if jc else ""
    paragraph_xml = f'<w:pPr xmlns:w="{W}">{jc_xml}</w:pPr>'
    converted = absolute_drawing_xml(
        [_inline_drawing_xml()],
        paragraph_xml=paragraph_xml,
        section_xml=_layout_section_xml(),
    )
    root = etree.fromstring(converted[0].encode())
    offset = root.find(f".//{{{WP}}}positionH/{{{WP}}}posOffset").text
    assert int(offset) == expected_twips * 635


def test_absolute_composite_horizontal_offset_includes_indents():
    paragraph_xml = f"""
      <w:pPr xmlns:w="{W}">
        <w:jc w:val="center"/>
        <w:ind w:left="100" w:right="50"/>
      </w:pPr>
    """
    converted = absolute_drawing_xml(
        [_inline_drawing_xml()],
        paragraph_xml=paragraph_xml,
        section_xml=_layout_section_xml(),
    )
    root = etree.fromstring(converted[0].encode())
    offset = root.find(f".//{{{WP}}}positionH/{{{WP}}}posOffset").text
    assert int(offset) == 225 * 635


def test_absolute_composite_horizontal_offset_clamps_overwide_image():
    paragraph_xml = f'<w:pPr xmlns:w="{W}"><w:jc w:val="right"/></w:pPr>'
    converted = absolute_drawing_xml(
        [_inline_drawing_xml(width=600000)],
        paragraph_xml=paragraph_xml,
        section_xml=_layout_section_xml(),
    )
    root = etree.fromstring(converted[0].encode())
    offset = root.find(f".//{{{WP}}}positionH/{{{WP}}}posOffset").text
    assert offset == "0"


def test_absolute_composite_shape_column_offset_includes_left_indent():
    paragraph_xml = f"""
      <w:pPr xmlns:w="{W}"><w:ind w:left="100"/></w:pPr>
    """
    drawing = f"""
      <w:drawing xmlns:w="{W}" xmlns:wp="{WP}">
        <wp:anchor>
          <wp:positionH relativeFrom="column"><wp:posOffset>20</wp:posOffset></wp:positionH>
          <wp:positionV relativeFrom="paragraph"><wp:posOffset>30</wp:posOffset></wp:positionV>
          <wp:extent cx="100" cy="100"/>
        </wp:anchor>
      </w:drawing>
    """
    converted = absolute_drawing_xml(
        [drawing],
        paragraph_xml=paragraph_xml,
        section_xml=_layout_section_xml(),
    )
    root = etree.fromstring(converted[0].encode())
    offset = root.find(f".//{{{WP}}}positionH/{{{WP}}}posOffset").text
    assert int(offset) == 100 * 635 + 20


def test_absolute_composite_vertical_offset_subtracts_half_grid_gap():
    paragraph_xml = f'<w:pPr xmlns:w="{W}"/>'
    section_xml = _layout_section_xml(
        f'<w:docGrid w:type="lines" w:linePitch="360"/>'
    )
    converted = absolute_drawing_xml(
        [_shape_drawing_xml(), _inline_drawing_xml(height=200000)],
        paragraph_xml=paragraph_xml,
        section_xml=section_xml,
    )
    root = etree.fromstring(converted[0].encode())
    offset = root.find(f".//{{{WP}}}positionV/{{{WP}}}posOffset").text
    expected_gap = (360 * 635 - (200000 % (360 * 635))) // 2
    assert int(offset) == 300000 - expected_gap


def test_absolute_composite_vertical_offset_keeps_exact_grid_height():
    paragraph_xml = f'<w:pPr xmlns:w="{W}"/>'
    section_xml = _layout_section_xml(
        f'<w:docGrid w:type="lineAndChar" w:linePitch="360"/>'
    )
    inline = _inline_drawing_xml().replace('cy="1000"', 'cy="228600"')
    converted = absolute_drawing_xml(
        [_shape_drawing_xml(), inline],
        paragraph_xml=paragraph_xml,
        section_xml=section_xml,
    )
    root = etree.fromstring(converted[0].encode())
    offset = root.find(f".//{{{WP}}}positionV/{{{WP}}}posOffset").text
    assert int(offset) == 300000


@pytest.mark.parametrize(
    "paragraph_xml, section_xml, drawings",
    [
        (
            f'<w:pPr xmlns:w="{W}"/>',
            _layout_section_xml(),
            [_shape_drawing_xml(), _inline_drawing_xml()],
        ),
        (
            f'<w:pPr xmlns:w="{W}"/>',
            _layout_section_xml(
                f'<w:docGrid w:type="default" w:linePitch="360"/>'
            ),
            [_shape_drawing_xml(), _inline_drawing_xml()],
        ),
        (
            f'<w:pPr xmlns:w="{W}"><w:snapToGrid w:val="0"/></w:pPr>',
            _layout_section_xml(
                f'<w:docGrid w:type="lines" w:linePitch="360"/>'
            ),
            [_shape_drawing_xml(), _inline_drawing_xml()],
        ),
        (
            f'<w:pPr xmlns:w="{W}"/>',
            _layout_section_xml(
                f'<w:docGrid w:type="lines" w:linePitch="360"/>'
            ),
            [_shape_drawing_xml()],
        ),
    ],
)
def test_absolute_composite_vertical_offset_skips_grid_gap_without_requirements(
    paragraph_xml, section_xml, drawings
):
    converted = absolute_drawing_xml(
        drawings,
        paragraph_xml=paragraph_xml,
        section_xml=section_xml,
    )
    root = etree.fromstring(converted[0].encode())
    offset = root.find(f".//{{{WP}}}positionV/{{{WP}}}posOffset").text
    assert int(offset) == 300000


def _numbering_xml():
    return f"""<w:numbering xmlns:w="{W}">
      <w:abstractNum w:abstractNumId="1">
        <w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1."/></w:lvl>
        <w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="lowerLetter"/><w:lvlText w:val="%1.%2."/></w:lvl>
      </w:abstractNum>
      <w:num w:numId="9"><w:abstractNumId w:val="1"/></w:num>
    </w:numbering>""".encode()


def _paragraph_with_num(level, number_id="9"):
    document = Document()
    paragraph = document.add_paragraph("項目")
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(level))
    num_id = OxmlElement("w:numId")
    num_id.set(qn("w:val"), number_id)
    num_pr.extend([ilvl, num_id])
    paragraph._p.get_or_add_pPr().append(num_pr)
    return paragraph


def test_numbering_resolver_tracks_nested_counters():
    resolver = NumberingResolver(_numbering_xml())
    assert resolver.marker(_paragraph_with_num(0))[1] == "1."
    assert resolver.marker(_paragraph_with_num(1))[1] == "1.a."
    assert resolver.marker(_paragraph_with_num(1))[1] == "1.b."
    assert resolver.marker(_paragraph_with_num(0))[1] == "2."
    assert resolver.marker(_paragraph_with_num(1))[1] == "2.a."


def test_numbering_resolver_honors_start_override_and_lvl_restart():
    blob = f"""<w:numbering xmlns:w="{W}">
      <w:abstractNum w:abstractNumId="2">
        <w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1."/></w:lvl>
        <w:lvl w:ilvl="1"><w:start w:val="1"/><w:lvlRestart w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1.%2."/></w:lvl>
      </w:abstractNum>
      <w:num w:numId="10"><w:abstractNumId w:val="2"/>
        <w:lvlOverride w:ilvl="0"><w:startOverride w:val="3"/></w:lvlOverride>
      </w:num>
    </w:numbering>""".encode()
    resolver = NumberingResolver(blob)
    assert resolver.marker(_paragraph_with_num(0, "10"))[1] == "3."
    assert resolver.marker(_paragraph_with_num(1, "10"))[1] == "3.1."
    assert resolver.marker(_paragraph_with_num(1, "10"))[1] == "3.2."
    paragraph = _paragraph_with_num(0, "10")
    assert resolver.marker(paragraph)[1] == "4."
    paragraph = _paragraph_with_num(1, "10")
    assert resolver.marker(paragraph)[1] == "4.1."


def test_numbering_resolver_zero_lvl_restart_preserves_lower_counter():
    blob = f"""<w:numbering xmlns:w="{W}">
      <w:abstractNum w:abstractNumId="20">
        <w:lvl w:ilvl="0"><w:numFmt w:val="decimal"/><w:lvlText w:val="%1."/></w:lvl>
        <w:lvl w:ilvl="1"><w:lvlRestart w:val="0"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1.%2."/></w:lvl>
      </w:abstractNum>
      <w:num w:numId="20"><w:abstractNumId w:val="20"/></w:num>
    </w:numbering>""".encode()
    resolver = NumberingResolver(blob)

    assert resolver.marker(_paragraph_with_num(0, "20"))[1] == "1."
    assert resolver.marker(_paragraph_with_num(1, "20"))[1] == "1.1."
    assert resolver.marker(_paragraph_with_num(0, "20"))[1] == "2."
    assert resolver.marker(_paragraph_with_num(1, "20"))[1] == "2.2."


def test_numbering_resolver_explicit_restart_level_controls_reset():
    blob = f"""<w:numbering xmlns:w="{W}">
      <w:abstractNum w:abstractNumId="21">
        <w:lvl w:ilvl="0"><w:numFmt w:val="decimal"/><w:lvlText w:val="%1."/></w:lvl>
        <w:lvl w:ilvl="1"><w:numFmt w:val="decimal"/><w:lvlText w:val="%1.%2."/></w:lvl>
        <w:lvl w:ilvl="2"><w:lvlRestart w:val="2"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1.%2.%3."/></w:lvl>
      </w:abstractNum>
      <w:num w:numId="21"><w:abstractNumId w:val="21"/></w:num>
    </w:numbering>""".encode()
    resolver = NumberingResolver(blob)

    assert resolver.marker(_paragraph_with_num(0, "21"))[1] == "1."
    assert resolver.marker(_paragraph_with_num(1, "21"))[1] == "1.1."
    assert resolver.marker(_paragraph_with_num(2, "21"))[1] == "1.1.1."
    assert resolver.marker(_paragraph_with_num(0, "21"))[1] == "2."
    assert resolver.marker(_paragraph_with_num(2, "21"))[1] == "2.1.2."
    assert resolver.marker(_paragraph_with_num(1, "21"))[1] == "2.1."
    assert resolver.marker(_paragraph_with_num(2, "21"))[1] == "2.1.1."


def test_safe_xml_parser_rejects_doctype_and_accepts_normal_xml():
    with pytest.raises(ET.ParseError, match="DOCTYPE"):
        safe_fromstring(
            b'<!DOCTYPE root [<!ENTITY value "expanded">]><root>&value;</root>'
        )
    assert safe_fromstring(b"<root />").tag == "root"


def test_numbering_resolver_skips_doctype_payload():
    blob = b"""<!DOCTYPE numbering [<!ENTITY value "expanded">]>
    <w:numbering xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
      <w:num w:numId="1"><w:abstractNumId w:val="1"/></w:num>
    </w:numbering>"""

    resolver = NumberingResolver(blob)

    assert not resolver.valid


def test_composite_xml_skips_doctype_payload():
    drawing = b"""<!DOCTYPE drawing [<!ENTITY value "expanded">]>
    <wp:inline xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" />"""

    assert absolute_drawing_xml([drawing]) is None


def test_mpp_error_message_uses_english_translation():
    i18n.setup_i18n("en")

    assert i18n._("MPP読み込みエラー: {error}").format(error="broken") == (
        "MPP reading error: broken"
    )

    i18n.setup_i18n("ja")


def test_unknown_field_result_is_kept_and_hyperlink_is_rendered():
    document = Document()
    paragraph = document.add_paragraph()
    begin = OxmlElement("w:r")
    begin_char = OxmlElement("w:fldChar")
    begin_char.set(qn("w:fldCharType"), "begin")
    begin.append(begin_char)
    instruction = OxmlElement("w:r")
    instr_text = OxmlElement("w:instrText")
    instr_text.text = 'HYPERLINK "https://example.test"'
    instruction.append(instr_text)
    separate = OxmlElement("w:r")
    separate_char = OxmlElement("w:fldChar")
    separate_char.set(qn("w:fldCharType"), "separate")
    separate.append(separate_char)
    result = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "リンク"
    result.append(text)
    end = OxmlElement("w:r")
    end_char = OxmlElement("w:fldChar")
    end_char.set(qn("w:fldCharType"), "end")
    end.append(end_char)
    paragraph._p.extend([begin, instruction, separate, result, end])
    assert convert_paragraph(paragraph) == "[リンク](https://example.test)"


@pytest.mark.parametrize(
    ("instruction_text", "expected"),
    [
        (
            'HYPERLINK "https://example.test" \\l "bookmark"',
            "[リンク](https://example.test#bookmark)",
        ),
        (
            'HYPERLINK \\l "bookmark"',
            "[リンク](#bookmark)",
        ),
        (
            'HYPERLINK "https://example.test" \\l',
            "[リンク](https://example.test)",
        ),
    ],
)
def test_hyperlink_field_combines_document_url_and_bookmark(
    instruction_text, expected
):
    document = Document()
    paragraph = document.add_paragraph()
    begin = OxmlElement("w:r")
    begin_char = OxmlElement("w:fldChar")
    begin_char.set(qn("w:fldCharType"), "begin")
    begin.append(begin_char)
    instruction = OxmlElement("w:r")
    instr_text = OxmlElement("w:instrText")
    instr_text.text = instruction_text
    instruction.append(instr_text)
    separate = OxmlElement("w:r")
    separate_char = OxmlElement("w:fldChar")
    separate_char.set(qn("w:fldCharType"), "separate")
    separate.append(separate_char)
    result = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "リンク"
    result.append(text)
    end = OxmlElement("w:r")
    end_char = OxmlElement("w:fldChar")
    end_char.set(qn("w:fldCharType"), "end")
    end.append(end_char)
    paragraph._p.extend([begin, instruction, separate, result, end])

    assert convert_paragraph(paragraph) == expected


@pytest.mark.parametrize(
    ("target", "expected"),
    [
        ("https://example.test/Shared%20Documents/file%20(1).html", "https://example.test/Shared%20Documents/file%20%281%29.html"),
        ("https://example.test/日本語/%E3%83%95%E3%82%A1%E3%82%A4%E3%83%AB", "https://example.test/日本語/ファイル"),
        ("https://example.test/100%done", "https://example.test/100%done"),
        ("#開始位置", "#開始位置"),
    ],
)
def test_markdown_url_normalizes_only_unsafe_delimiters(target, expected):
    assert normalize_markdown_url(target) == expected


def test_hyperlink_field_escapes_markdown_url_delimiters():
    document = Document()
    paragraph = document.add_paragraph()
    begin = OxmlElement("w:r")
    begin_char = OxmlElement("w:fldChar")
    begin_char.set(qn("w:fldCharType"), "begin")
    begin.append(begin_char)
    instruction = OxmlElement("w:r")
    instr_text = OxmlElement("w:instrText")
    instr_text.text = 'HYPERLINK "https://example.test/Shared%20Documents/file%20(1).html"'
    instruction.append(instr_text)
    separate = OxmlElement("w:r")
    separate_char = OxmlElement("w:fldChar")
    separate_char.set(qn("w:fldCharType"), "separate")
    separate.append(separate_char)
    result = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "リンク"
    result.append(text)
    end = OxmlElement("w:r")
    end_char = OxmlElement("w:fldChar")
    end_char.set(qn("w:fldCharType"), "end")
    end.append(end_char)
    paragraph._p.extend([begin, instruction, separate, result, end])

    assert convert_paragraph(paragraph) == (
        "[リンク](https://example.test/Shared%20Documents/file%20%281%29.html)"
    )


def test_hyperlink_element_normalizes_external_and_internal_targets():
    document = Document()
    paragraph = document.add_paragraph()
    external = OxmlElement("w:hyperlink")
    external.set(
        "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id",
        "rId9",
    )
    external_run = OxmlElement("w:r")
    external_text = OxmlElement("w:t")
    external_text.text = "外部リンク"
    external_run.append(external_text)
    external.append(external_run)
    internal = OxmlElement("w:hyperlink")
    internal.set(qn("w:anchor"), "開始 位置(1)")
    internal_run = OxmlElement("w:r")
    internal_text = OxmlElement("w:t")
    internal_text.text = "内部リンク"
    internal_run.append(internal_text)
    internal.append(internal_run)
    paragraph._p.extend([external, internal])

    assert convert_paragraph(
        paragraph,
        hyperlink_resolver={"rId9": "https://example.test/Shared Documents"}.get,
    ) == (
        "[外部リンク](https://example.test/Shared%20Documents)"
        "[内部リンク](#開始%20位置%281%29)"
    )


def test_field_literals_and_adjacent_simple_fields_are_preserved():
    document = Document()
    paragraph = document.add_paragraph("前")

    first = OxmlElement("w:fldSimple")
    first.set(qn("w:instr"), " STYLEREF 1 \\s ")
    first_result = OxmlElement("w:r")
    first_text = OxmlElement("w:t")
    first_text.text = "9"
    first_result.append(first_text)
    first.append(first_result)

    separator = OxmlElement("w:r")
    separator.append(OxmlElement("w:noBreakHyphen"))

    second = OxmlElement("w:fldSimple")
    second.set(qn("w:instr"), " SEQ 図 \\* ARABIC ")
    second_result = OxmlElement("w:r")
    second_text = OxmlElement("w:t")
    second_text.text = "1"
    second_result.append(second_text)
    second.append(second_result)

    paragraph._p.extend([first, separator, second])
    paragraph.add_run("後")

    assert convert_paragraph(paragraph) == "前9-1後"


def test_hyperlink_element_resolves_external_and_internal_targets():
    document = Document()
    paragraph = document.add_paragraph()
    external = OxmlElement("w:hyperlink")
    external.set("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id", "rId9")
    external_run = OxmlElement("w:r")
    external_text = OxmlElement("w:t")
    external_text.text = "外部リンク"
    external_run.append(external_text)
    external.append(external_run)
    internal = OxmlElement("w:hyperlink")
    internal.set(qn("w:anchor"), "開始位置")
    internal_run = OxmlElement("w:r")
    internal_text = OxmlElement("w:t")
    internal_text.text = "内部リンク"
    internal_run.append(internal_text)
    internal.append(internal_run)
    paragraph._p.append(external)
    paragraph._p.append(internal)

    assert convert_paragraph(
        paragraph,
        hyperlink_resolver={"rId9": "https://example.test"}.get,
    ) == "[外部リンク](https://example.test)[内部リンク](#開始位置)"


def test_hyperlink_unresolved_target_keeps_text_and_nested_containers():
    document = Document()
    paragraph = document.add_paragraph()
    sdt = OxmlElement("w:sdt")
    content = OxmlElement("w:sdtContent")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "content control"
    run.append(text)
    content.append(run)
    sdt.append(content)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink_run = OxmlElement("w:r")
    hyperlink_text = OxmlElement("w:t")
    hyperlink_text.text = "リンク文字"
    hyperlink_run.append(hyperlink_text)
    hyperlink.append(hyperlink_run)
    paragraph._p.extend([sdt, hyperlink])

    assert convert_paragraph(paragraph) == "content controlリンク文字"


def test_top_level_sdt_recursively_converts_paragraph_and_table(tmp_path):
    document = Document()
    document.add_paragraph("前段落")
    sdt = OxmlElement("w:sdt")
    properties = OxmlElement("w:sdtPr")
    content = OxmlElement("w:sdtContent")
    nested = OxmlElement("w:sdt")
    nested_content = OxmlElement("w:sdtContent")
    paragraph = OxmlElement("w:p")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "SDT内段落"
    run.append(text)
    paragraph.append(run)
    nested_content.append(paragraph)
    nested.append(nested_content)
    content.append(nested)
    table = OxmlElement("w:tbl")
    table_properties = OxmlElement("w:tblPr")
    table_grid = OxmlElement("w:tblGrid")
    grid_column = OxmlElement("w:gridCol")
    grid_column.set(qn("w:w"), "1000")
    table_grid.append(grid_column)
    row = OxmlElement("w:tr")
    cell = OxmlElement("w:tc")
    cell_properties = OxmlElement("w:tcPr")
    cell_width = OxmlElement("w:tcW")
    cell_width.set(qn("w:w"), "1000")
    cell_width.set(qn("w:type"), "dxa")
    cell_properties.append(cell_width)
    cell_paragraph = OxmlElement("w:p")
    cell_run = OxmlElement("w:r")
    cell_text = OxmlElement("w:t")
    cell_text.text = "SDT内表"
    cell_run.append(cell_text)
    cell_paragraph.append(cell_run)
    cell.append(cell_properties)
    cell.append(cell_paragraph)
    row.append(cell)
    table.extend([table_properties, table_grid, row])
    content.append(table)
    sdt.extend([properties, content])
    document._body._element.append(sdt)
    path = tmp_path / "top-level-sdt.docx"
    document.save(path)

    from o2md.d2md import WordToMarkdownConverter

    output_path = WordToMarkdownConverter(str(path), output_dir=str(tmp_path)).convert()
    output = Path(output_path).read_text()
    assert "前段落" in output
    assert "SDT内段落" in output
    assert "| SDT内表 |" in output
    assert output.count("SDT内段落") == 1


def test_top_level_toc_sdt_does_not_duplicate_generated_toc(tmp_path):
    document = Document()
    document.add_paragraph("目次")
    document.add_heading("見出し", level=1)
    sdt = OxmlElement("w:sdt")
    properties = OxmlElement("w:sdtPr")
    gallery = OxmlElement("w:docPartObj")
    gallery_name = OxmlElement("w:docPartGallery")
    gallery_name.set(qn("w:val"), "Table of Contents")
    gallery.append(gallery_name)
    properties.append(gallery)
    content = OxmlElement("w:sdtContent")
    paragraph = OxmlElement("w:p")
    run = OxmlElement("w:r")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.text = ' TOC \\o "1-3" '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run.extend([begin, instruction, separate, end])
    paragraph.append(run)
    content.append(paragraph)
    sdt.extend([properties, content])
    document._body._element.append(sdt)
    path = tmp_path / "toc-sdt.docx"
    document.save(path)

    from o2md.d2md import WordToMarkdownConverter

    output_path = WordToMarkdownConverter(str(path), output_dir=str(tmp_path)).convert()
    output = Path(output_path).read_text()
    assert output.count("# 目次") == 1
    assert "TOC" not in output


def test_empty_note_reference_is_omitted_but_surrounding_text_kept(tmp_path):
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("前")
    reference_run = paragraph.add_run()
    reference = OxmlElement("w:footnoteReference")
    reference.set(qn("w:id"), "1")
    reference_run._r.append(reference)
    paragraph.add_run("後")
    path = tmp_path / "empty-note.docx"
    document.save(path)
    rebuilt = tmp_path / "empty-note-rebuilt.docx"
    footnotes = f"""<w:footnotes xmlns:w="{W}">
      <w:footnote w:id="-1" w:type="separator"/>
      <w:footnote w:id="1"><w:p/></w:footnote>
    </w:footnotes>""".encode()
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(rebuilt, "w") as target:
        for item in source.infolist():
            if item.filename == "word/footnotes.xml":
                target.writestr(item, footnotes)
            else:
                target.writestr(item, source.read(item.filename))

    from o2md.d2md import WordToMarkdownConverter

    output_path = WordToMarkdownConverter(
        str(rebuilt), output_dir=str(tmp_path)
    ).convert()
    output = Path(output_path).read_text()
    assert "前後" in output
    assert "[^fn1]" not in output


def test_hyperlink_inside_field_keeps_link_text_and_markup():
    document = Document()
    paragraph = document.add_paragraph()
    begin = OxmlElement("w:r")
    begin_char = OxmlElement("w:fldChar")
    begin_char.set(qn("w:fldCharType"), "begin")
    begin.append(begin_char)
    instruction = OxmlElement("w:r")
    instr_text = OxmlElement("w:instrText")
    instr_text.text = " REF target "
    instruction.append(instr_text)
    separate = OxmlElement("w:r")
    separate_char = OxmlElement("w:fldChar")
    separate_char.set(qn("w:fldCharType"), "separate")
    separate.append(separate_char)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), "target")
    link_run = OxmlElement("w:r")
    link_text = OxmlElement("w:t")
    link_text.text = "リンク"
    link_run.append(link_text)
    hyperlink.append(link_run)
    end = OxmlElement("w:r")
    end_char = OxmlElement("w:fldChar")
    end_char.set(qn("w:fldCharType"), "end")
    end.append(end_char)
    paragraph._p.extend([begin, instruction, separate, hyperlink, end])

    assert convert_paragraph(paragraph) == "[リンク](#target)"


def test_numbering_resolver_uses_start_for_unseen_parent_level():
    blob = f"""<w:numbering xmlns:w="{W}">
      <w:abstractNum w:abstractNumId="3">
        <w:lvl w:ilvl="0"><w:start w:val="4"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1."/></w:lvl>
        <w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="lowerLetter"/><w:lvlText w:val="%1.%2."/></w:lvl>
      </w:abstractNum>
      <w:num w:numId="13"><w:abstractNumId w:val="3"/></w:num>
    </w:numbering>""".encode()
    resolver = NumberingResolver(blob)

    assert resolver.marker(_paragraph_with_num(1, "13"))[1] == "4.a."


def test_numbering_resolver_normalizes_simple_non_decimal_labels():
    blob = f"""<w:numbering xmlns:w="{W}">
      <w:abstractNum w:abstractNumId="4">
        <w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="upperRoman"/><w:lvlText w:val="%1."/></w:lvl>
        <w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="decimalEnclosedCircle"/><w:lvlText w:val="%2)"/></w:lvl>
        <w:lvl w:ilvl="2"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="第%3章"/></w:lvl>
        <w:lvl w:ilvl="3"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="・"/></w:lvl>
      </w:abstractNum>
      <w:num w:numId="14"><w:abstractNumId w:val="4"/></w:num>
    </w:numbering>""".encode()
    resolver = NumberingResolver(blob)

    assert resolver.marker(_paragraph_with_num(0, "14"))[1] == "1."
    assert resolver.marker(_paragraph_with_num(1, "14"))[1] == "1."
    assert resolver.marker(_paragraph_with_num(2, "14"))[1] == "第1章"
    assert resolver.marker(_paragraph_with_num(3, "14"))[1:] == ("-", True)


def test_field_paragraph_keeps_bold_run_with_note_reference(tmp_path):
    path = tmp_path / "source.docx"
    document = Document()
    paragraph = document.add_paragraph()
    run = paragraph.add_run("太字")
    run.bold = True
    reference_run = paragraph.add_run()
    reference = OxmlElement("w:footnoteReference")
    reference.set(qn("w:id"), "1")
    reference_run._r.append(reference)
    document.save(path)
    rebuilt = tmp_path / "with-note.docx"
    footnotes = f"""<w:footnotes xmlns:w="{W}">
      <w:footnote w:id="1"><w:p><w:r><w:t>脚注本文</w:t></w:r></w:p></w:footnote>
    </w:footnotes>""".encode()
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(rebuilt, "w") as target:
        for item in source.infolist():
            target.writestr(item, source.read(item.filename))
        target.writestr("word/footnotes.xml", footnotes)
    from o2md.d2md import WordToMarkdownConverter

    result = WordToMarkdownConverter(str(rebuilt), output_dir=str(tmp_path)).convert()
    output = Path(result).read_text()
    assert "**太字**[^fn1]" in output


def test_broken_optional_docx_part_keeps_body(tmp_path):
    path = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("本文")
    document.save(path)
    rebuilt = tmp_path / "broken.docx"
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(rebuilt, "w") as target:
        for item in source.infolist():
            if item.filename == "word/numbering.xml":
                target.writestr(item, b"<broken")
            else:
                target.writestr(item, source.read(item.filename))
    from o2md.d2md import WordToMarkdownConverter

    result = WordToMarkdownConverter(str(rebuilt), output_dir=str(tmp_path)).convert()
    assert "本文" in Path(result).read_text()


def test_broken_docx_and_degradation_failure_preserves_cause(tmp_path):
    path = tmp_path / "broken.docx"
    path.write_bytes(b"not a zip archive")
    from o2md.d2md import WordToMarkdownConverter

    converter = WordToMarkdownConverter.__new__(WordToMarkdownConverter)
    with pytest.raises(RuntimeError, match="DOCX本文の読み込みに失敗しました") as error:
        converter._load_document_with_optional_degradation(str(path))
    assert error.value.__cause__ is not None
    assert not isinstance(error.value.__cause__, NameError)


def test_table_cell_formatting_preserves_runs_and_fields(tmp_path):
    path = tmp_path / "table-cell.docx"
    document = Document()
    table = document.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    paragraph = cell.paragraphs[0]
    bold = paragraph.add_run("太字")
    bold.bold = True
    hidden = paragraph.add_run("隠し文字")
    hidden._r.get_or_add_rPr().append(OxmlElement("w:vanish"))
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), " DATE ")
    result = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "フィールド結果"
    result.append(text)
    field.append(result)
    paragraph._p.append(field)
    document.save(path)

    from o2md.d2md import WordToMarkdownConverter

    output_path = WordToMarkdownConverter(
        str(path), output_dir=str(tmp_path)
    ).convert()
    output = Path(output_path).read_text()
    assert "**太字**" in output
    assert "隠し文字" not in output
    assert "フィールド結果" in output


def test_table_cell_line_breaks_are_markdown_safe(tmp_path):
    path = tmp_path / "table-cell-breaks.docx"
    document = Document()
    table = document.add_table(rows=1, cols=1)
    paragraph = table.cell(0, 0).paragraphs[0]
    paragraph.add_run("行内")
    paragraph.add_run().add_break()
    paragraph.add_run("改行")
    table.cell(0, 0).add_paragraph("段落")
    document.save(path)

    from o2md.d2md import WordToMarkdownConverter

    output_path = WordToMarkdownConverter(
        str(path), output_dir=str(tmp_path)
    ).convert()
    output = Path(output_path).read_text()
    assert "| 行内<br>改行<br>段落 |" in output
    assert "行内\n改行" not in output


def test_table_keeps_legacy_spacing_after_table(tmp_path):
    path = tmp_path / "table-spacing.docx"
    document = Document()
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "表"
    document.add_paragraph("後続段落")
    document.save(path)

    from o2md.d2md import WordToMarkdownConverter

    output_path = WordToMarkdownConverter(
        str(path), output_dir=str(tmp_path)
    ).convert()
    output = Path(output_path).read_text()
    assert "| 表 |" in output
    assert "| 表 |\n| --- |\n\n\n後続段落" in output


def test_mspdi_content_sniffing_is_not_extension_only():
    data = f'<Project xmlns="{W.replace("openxmlformats.org/wordprocessingml/2006/main", "schemas.microsoft.com/project")}"/>'.encode()
    assert is_mspdi_xml(data)
    assert detect_type_from_bytes(data) == "msproject"
    assert detect_type_from_bytes(b"<Project xmlns='http://example.test'/>") == "unknown"


def test_mspdi_xml_rejects_doctype_without_entity_expansion():
    data = b"""<!DOCTYPE lolz [
      <!ENTITY a "1234567890">
      <!ENTITY b "&a;&a;&a;&a;&a;&a;&a;&a;&a;&a;">
    ]>
    <Project xmlns="http://schemas.microsoft.com/project">&b;</Project>"""

    assert not is_mspdi_xml(data)


def test_notes_manager_skips_doctype_note_part(tmp_path):
    document = Document()
    path = tmp_path / "notes-dtd.docx"
    document.save(path)
    footnotes = f"""<!DOCTYPE footnotes [
      <!ENTITY text "危険な脚注">
    ]>
    <w:footnotes xmlns:w="{W}">
      <w:footnote w:id="1"><w:p><w:r><w:t>&text;</w:t></w:r></w:p></w:footnote>
    </w:footnotes>""".encode()
    rebuilt = tmp_path / "notes-dtd-rebuilt.docx"
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(rebuilt, "w") as target:
        for item in source.infolist():
            target.writestr(item, source.read(item.filename))
        target.writestr("word/footnotes.xml", footnotes)

    manager = NoteManager(str(rebuilt))

    assert manager.reference("fn", "1") == ""


def test_notes_manager_maps_references_and_definitions(tmp_path):
    document = Document()
    path = tmp_path / "notes.docx"
    document.save(path)
    footnotes = f"""<w:footnotes xmlns:w="{W}">
      <w:footnote w:id="-1" w:type="separator"/>
      <w:footnote w:id="1"><w:p><w:r><w:t>脚注本文</w:t></w:r></w:p></w:footnote>
    </w:footnotes>""".encode()
    rebuilt = tmp_path / "notes-with-footnote.docx"
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(rebuilt, "w") as target:
        for item in source.infolist():
            target.writestr(item, source.read(item.filename))
        target.writestr("word/footnotes.xml", footnotes)
    manager = NoteManager(str(rebuilt))
    assert manager.reference("fn", "1") == "[^fn1]"
    assert manager.definitions() == ["[^fn1]: 脚注本文"]


def test_notes_manager_definitions_handles_nested_note_reference(tmp_path):
    document = Document()
    path = tmp_path / "nested-notes.docx"
    document.save(path)
    footnotes = f"""<w:footnotes xmlns:w="{W}">
      <w:footnote w:id="1"><w:p><w:r><w:t>1つ目の脚注</w:t></w:r></w:p></w:footnote>
      <w:footnote w:id="2"><w:p><w:r><w:t>2つ目の脚注</w:t></w:r></w:p></w:footnote>
    </w:footnotes>""".encode()
    rebuilt = tmp_path / "nested-notes-rebuilt.docx"
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(rebuilt, "w") as target:
        for item in source.infolist():
            target.writestr(item, source.read(item.filename))
        target.writestr("word/footnotes.xml", footnotes)
    manager = NoteManager(str(rebuilt))
    assert manager.reference("fn", "1") == "[^fn1]"

    def renderer(node):
        text = "".join(t.text for t in node.iter(f"{{{W}}}t") if t.text)
        # 脚注本文の変換中に別の脚注参照が現れる状況を再現する。
        if text == "1つ目の脚注":
            text += manager.reference("fn", "2")
        return text

    definitions = manager.definitions(renderer=renderer)

    assert definitions == ["[^fn1]: 1つ目の脚注[^fn2]", "[^fn2]: 2つ目の脚注"]


def test_word_table_grid_keeps_column_count():
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "見出し"
    table.cell(0, 1).text = "列2"
    table.cell(1, 0).text = "値"
    lines = render_table(table, lambda cell: cell.text)
    assert lines[0] == "| 見出し | 列2 |"
    assert all(line.count("|") == 3 for line in lines)


def test_word_table_grid_without_tbl_grid_keeps_columns():
    document = Document()
    table = document.add_table(rows=2, cols=3)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    table.cell(0, 2).text = "C"
    table._tbl.remove(table._tbl.tblGrid)
    lines = render_table(table, lambda cell: cell.text)
    assert all(line.count("|") == 4 for line in lines)


def test_word_table_complex_merges_and_headers_keep_grid():
    document = Document()
    table = document.add_table(rows=5, cols=4)
    for row_index, row in enumerate(table.rows):
        for col_index, cell in enumerate(row.cells):
            cell.text = f"{row_index}-{col_index}"
    rows = table._tbl.tr_lst
    for row in rows[:2]:
        tr_pr = row.get_or_add_trPr()
        header = OxmlElement("w:tblHeader")
        tr_pr.append(header)
    first = rows[0].tc_lst[0]
    first.get_or_add_tcPr().append(OxmlElement("w:gridSpan"))
    first.tcPr.gridSpan.set(qn("w:val"), "2")
    rows[0].remove(rows[0].tc_lst[1])
    row1_pr = rows[1].get_or_add_trPr()
    before = OxmlElement("w:gridBefore")
    before.set(qn("w:val"), "1")
    after = OxmlElement("w:gridAfter")
    after.set(qn("w:val"), "1")
    row1_pr.extend([before, after])
    for row in (rows[2], rows[3], rows[4]):
        merge = OxmlElement("w:vMerge")
        merge.set(qn("w:val"), "restart" if row is rows[2] else "continue")
        row.tc_lst[0].get_or_add_tcPr().append(merge)
    span = OxmlElement("w:gridSpan")
    span.set(qn("w:val"), "2")
    rows[3].tc_lst[1].get_or_add_tcPr().append(span)
    rows[3].remove(rows[3].tc_lst[2])
    lines = render_table(table, lambda cell: cell.text)
    assert all(line.count("|") == 5 for line in lines)
    assert "0-0" in lines[0]
    assert "0-2" in lines[0]


def test_word_table_vmerge_does_not_shift_following_unmerged_row():
    document = Document()
    table = document.add_table(rows=3, cols=3)
    for row_index, row in enumerate(table.rows):
        for col_index, cell in enumerate(row.cells):
            cell.text = f"{row_index}-{col_index}"
    first = table._tbl.tr_lst[0].tc_lst[0]
    merge = OxmlElement("w:vMerge")
    merge.set(qn("w:val"), "restart")
    first.get_or_add_tcPr().append(merge)

    lines = render_table(table, lambda cell: cell.text)

    assert lines[0] == "| 0-0 | 0-1 | 0-2 |"
    assert lines[2] == "| 1-0 | 1-1 | 1-2 |"
    assert lines[3] == "| 2-0 | 2-1 | 2-2 |"


def test_word_table_hmerge_uses_each_row_cell_declaration():
    document = Document()
    table = document.add_table(rows=2, cols=3)
    for row_index, row in enumerate(table.rows):
        for col_index, cell in enumerate(row.cells):
            cell.text = f"{row_index}-{col_index}"
    cells = table._tbl.tr_lst[0].tc_lst
    restart = OxmlElement("w:hMerge")
    restart.set(qn("w:val"), "restart")
    cells[0].get_or_add_tcPr().append(restart)
    continuation = OxmlElement("w:hMerge")
    continuation.set(qn("w:val"), "continue")
    cells[1].get_or_add_tcPr().append(continuation)

    lines = render_table(table, lambda cell: cell.text)

    assert lines[0] == "| 0-0 | 0-0 | 0-2 |"
    assert lines[2] == "| 1-0 | 1-1 | 1-2 |"


def test_word_table_gridspan_repeats_origin_label_and_keeps_grid():
    document = Document()
    table = document.add_table(rows=2, cols=3)
    for row_index, row in enumerate(table.rows):
        for col_index, cell in enumerate(row.cells):
            cell.text = f"{row_index}-{col_index}"
    first = table._tbl.tr_lst[0].tc_lst[0]
    span = OxmlElement("w:gridSpan")
    span.set(qn("w:val"), "2")
    first.get_or_add_tcPr().append(span)
    table._tbl.tr_lst[0].remove(table._tbl.tr_lst[0].tc_lst[1])

    lines = render_table(table, lambda cell: cell.text)

    assert lines[0] == "| 0-0 | 0-0 | 0-2 |"
    assert all(line.count("|") == 4 for line in lines)


def test_word_table_three_row_vmerge_repeats_origin_label_and_keeps_grid():
    document = Document()
    table = document.add_table(rows=3, cols=2)
    for row_index, row in enumerate(table.rows):
        for col_index, cell in enumerate(row.cells):
            cell.text = f"{row_index}-{col_index}"
    rows = table._tbl.tr_lst
    for index, row in enumerate(rows):
        merge = OxmlElement("w:vMerge")
        merge.set(qn("w:val"), "restart" if index == 0 else "continue")
        row.tc_lst[0].get_or_add_tcPr().append(merge)

    lines = render_table(table, lambda cell: cell.text)

    assert lines[0] == "| 0-0 | 0-1 |"
    assert lines[2] == "| 0-0 | 1-1 |"
    assert lines[3] == "| 0-0 | 2-1 |"
    assert all(line.count("|") == 3 for line in lines)


def test_word_table_header_inherits_horizontal_merge_labels():
    document = Document()
    table = document.add_table(rows=2, cols=4)
    values = [
        ("見出しA", "横結合対象", "見出しC", "見出しD"),
        ("補助見出しA", "補助見出しB", "補助見出しC", "補助見出しD"),
    ]
    for row, values_row in zip(table.rows, values):
        for cell, value in zip(row.cells, values_row):
            cell.text = value
    rows = table._tbl.tr_lst
    for row in rows:
        tr_pr = row.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:tblHeader"))
    table._tbl.tblGrid.append(OxmlElement("w:gridCol"))
    first = rows[0].tc_lst[0]
    first.get_or_add_tcPr().append(OxmlElement("w:gridSpan"))
    first.tcPr.gridSpan.set(qn("w:val"), "2")
    rows[0].remove(rows[0].tc_lst[1])
    tr_pr = rows[1].get_or_add_trPr()
    before = OxmlElement("w:gridBefore")
    before.set(qn("w:val"), "1")
    after = OxmlElement("w:gridAfter")
    after.set(qn("w:val"), "0")
    tr_pr.extend([before, after])

    lines = render_table(table, lambda cell: cell.text)

    assert lines[0] == (
        "| 見出しA | 見出しA 補助見出しA | 見出しC 補助見出しB | "
        "見出しD 補助見出しC | 補助見出しD |"
    )


def test_word_table_header_deduplicates_vertical_merge_labels():
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "縦結合見出し"
    table.cell(0, 1).text = "上段見出し"
    table.cell(1, 0).text = "縦結合見出し"
    table.cell(1, 1).text = "下段見出し"
    rows = table._tbl.tr_lst
    for row in rows:
        row.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    merge = OxmlElement("w:vMerge")
    merge.set(qn("w:val"), "restart")
    rows[0].tc_lst[0].get_or_add_tcPr().append(merge)
    merge = OxmlElement("w:vMerge")
    merge.set(qn("w:val"), "continue")
    rows[1].tc_lst[0].get_or_add_tcPr().append(merge)

    lines = render_table(table, lambda cell: cell.text)

    assert lines[0] == "| 縦結合見出し | 上段見出し 下段見出し |"
    assert all(line.count("|") == 3 for line in lines)


def test_word_table_header_keeps_consecutive_distinct_labels():
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "上段ラベル"
    table.cell(0, 1).text = "上段列2"
    table.cell(1, 0).text = "下段ラベル"
    table.cell(1, 1).text = "下段列2"
    for row in table._tbl.tr_lst:
        row.get_or_add_trPr().append(OxmlElement("w:tblHeader"))

    lines = render_table(table, lambda cell: cell.text)

    assert lines[0] == "| 上段ラベル 下段ラベル | 上段列2 下段列2 |"
    assert all(line.count("|") == 3 for line in lines)


def test_word_table_three_headers_skip_empty_middle_cells():
    document = Document()
    table = document.add_table(rows=3, cols=2)
    values = [
        ("上段A", "上段B"),
        ("中段A", ""),
        ("下段A", "下段B"),
    ]
    for row, values_row in zip(table.rows, values):
        for cell, value in zip(row.cells, values_row):
            cell.text = value
        row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))

    lines = render_table(table, lambda cell: cell.text)

    assert lines[0] == "| 上段A 中段A 下段A | 上段B 下段B |"


def test_excel_sheet_failure_isolated_and_all_failure_raises(tmp_path):
    from openpyxl import Workbook

    path = tmp_path / "sheets.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "first"
    workbook.create_sheet("Second")["A1"] = "second"
    workbook.save(path)
    converter = ExcelToMarkdownConverter(str(path), output_dir=str(tmp_path))
    original = converter._convert_sheet
    calls = {"count": 0}

    def fail_first(sheet):
        calls["count"] += 1
        if calls["count"] == 1:
            raise RuntimeError("壊れたシート")
        return original(sheet)

    converter._convert_sheet = fail_first
    output = Path(converter.convert()).read_text()
    assert "second" in output

    converter = ExcelToMarkdownConverter(str(path), output_dir=str(tmp_path))
    converter._convert_sheet = lambda sheet: (_ for _ in ()).throw(RuntimeError("壊れたシート"))
    with pytest.raises(RuntimeError, match="すべてのシート"):
        converter.convert()


def test_mspdi_file_path_detection_and_plain_xml(tmp_path):
    project = tmp_path / "project.xml"
    project.write_text('<Project xmlns="http://schemas.microsoft.com/project"/>')
    plain = tmp_path / "plain.xml"
    plain.write_text("<root/>")
    assert detect_file_type(str(project)) == "msproject"
    assert detect_file_type(str(plain)) == "unknown"


def test_filter_file_and_stdin_routes_sniff_mspdi(tmp_path, monkeypatch, capsys):
    project = tmp_path / "project.xml"
    data = b'<Project xmlns="http://schemas.microsoft.com/project"/>'
    project.write_bytes(data)
    monkeypatch.setattr(filter_cli, "filter_file", lambda path, **kwargs: "変換結果")

    monkeypatch.setattr(sys, "argv", ["o2md-filter", str(project)])
    filter_cli.main()
    assert capsys.readouterr().out == "変換結果"

    monkeypatch.setattr(sys, "argv", ["o2md-filter"])
    monkeypatch.setattr(sys, "stdin", io.TextIOWrapper(io.BytesIO(data)))
    filter_cli.main()
    assert capsys.readouterr().out == "変換結果"


@pytest.mark.parametrize(
    ("name", "payload", "expected_suffix"),
    [
        (
            "project-without-extension",
            b'<Project xmlns="http://schemas.microsoft.com/project"/>',
            ".xml",
        ),
        (
            "document-without-extension",
            b"PK\x03\x04",
            ".docx",
        ),
    ],
)
def test_filter_file_sniffed_type_uses_converter_extension(
    tmp_path, monkeypatch, name, payload, expected_suffix
):
    if expected_suffix == ".docx":
        archive = io.BytesIO()
        with zipfile.ZipFile(archive, "w") as package:
            package.writestr(
                "[Content_Types].xml",
                '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                '<Override PartName="/word/document.xml" '
                'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
                "</Types>",
            )
        payload = archive.getvalue()
    source = tmp_path / name
    source.write_bytes(payload)
    captured = {}

    def fake_convert(path, output_dir=None, **_kwargs):
        captured["path"] = path
        output = Path(output_dir) / "result.txt"
        output.write_text("変換結果", encoding="utf-8")
        return str(output), {}, 0

    monkeypatch.setattr(
        "o2md.o2md.convert_office_to_markdown", fake_convert
    )
    result = filter_cli.filter_file(str(source))

    assert result == "変換結果"
    assert Path(captured["path"]).suffix == expected_suffix
    assert not Path(captured["path"]).exists()


def test_resource_duration_totals_use_preconverted_days():
    resources = [{"id": 1, "name": "担当者"}]
    tasks = [
        {"resources": "担当者", "duration": "8時間", "duration_days": 1.0},
        {"resources": "担当者", "duration": "2週", "duration_days": 10.0},
        {"resources": "担当者", "duration": "1ヶ月", "duration_days": 20.0},
        {"resources": "担当者", "duration": "3日", "duration_days": 3.0},
    ]

    output = resources_to_markdown_table(resources, tasks)

    assert "| 1 | 担当者 | 4 | 34日 |" in output


def test_resource_duration_totals_note_unconvertible_values():
    resources = [{"id": 1, "name": "担当者"}]
    tasks = [
        {"resources": "担当者", "duration": "8時間", "duration_days": 1.0},
        {"resources": "担当者", "duration": "不明", "duration_days": None},
    ]

    output = resources_to_markdown_table(resources, tasks)

    assert "1日（1件は換算不能）" in output
